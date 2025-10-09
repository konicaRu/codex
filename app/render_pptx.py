from __future__ import annotations

from pathlib import Path
from typing import Iterable

from pptx import Presentation
from pptx.util import Inches, Pt

from .models import ActionsResult, MinutesResult, TopicsResult, VideoSegment


class PPTXRenderer:
    def __init__(self, output_dir: Path) -> None:
        self.output_dir = output_dir
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def render(
        self,
        video_name: str,
        segments: Iterable[VideoSegment],
        topics: TopicsResult | None,
        actions: ActionsResult | None,
        minutes: MinutesResult | None,
    ) -> Path:
        output_path = self.output_dir / f"{video_name}_report.pptx"
        presentation = Presentation()
        presentation.slide_width = Inches(13.33)
        presentation.slide_height = Inches(7.5)

        for segment in segments:
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12.5), Inches(1.0))
            title_tf = title_box.text_frame
            title_tf.text = segment.summary_title or "Сегмент"
            title_tf.paragraphs[0].font.size = Pt(28)
            title_tf.paragraphs[0].font.bold = True

            bullet_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.4), Inches(6.5), Inches(5))
            bullet_tf = bullet_box.text_frame
            bullet_tf.word_wrap = True
            for idx, bullet in enumerate(segment.summary_bullets[:3]):
                if idx == 0:
                    bullet_tf.text = bullet
                else:
                    p = bullet_tf.add_paragraph()
                    p.text = bullet
                bullet_tf.paragraphs[idx].font.size = Pt(18)

            if segment.best_frame_path.exists():
                slide.shapes.add_picture(
                    str(segment.best_frame_path),
                    Inches(7.2),
                    Inches(1.4),
                    width=Inches(5.8),
                )

        self._add_summary_slide(presentation, topics, actions, minutes)
        presentation.save(str(output_path))
        return output_path

    def _add_summary_slide(
        self,
        presentation: Presentation,
        topics: TopicsResult | None,
        actions: ActionsResult | None,
        minutes: MinutesResult | None,
    ) -> None:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12.5), Inches(1))
        title_tf = title_box.text_frame
        title_tf.text = "Итоги встречи"
        title_tf.paragraphs[0].font.size = Pt(32)
        title_tf.paragraphs[0].font.bold = True

        y = Inches(1.4)
        if topics:
            y = self._write_section(slide, "Темы", topics.topics, y)
            if topics.summary:
                y = self._write_section(slide, "Резюме", [topics.summary], y)

        if actions and actions.actions:
            y = self._write_section(slide, "Экшен-пункты", actions.actions, y)

        if minutes and minutes.bullets:
            self._write_section(slide, "Главные выводы", minutes.bullets, y)

    def _write_section(
        self,
        slide,
        title: str,
        lines: Iterable[str],
        top: float,
    ) -> float:
        section_title = slide.shapes.add_textbox(Inches(0.5), top, Inches(12.5), Inches(0.5))
        title_tf = section_title.text_frame
        title_tf.text = title
        title_tf.paragraphs[0].font.size = Pt(24)
        title_tf.paragraphs[0].font.bold = True

        bullets_box = slide.shapes.add_textbox(Inches(0.5), top + Inches(0.6), Inches(12.5), Inches(1.5))
        bullets_tf = bullets_box.text_frame
        bullets_tf.word_wrap = True
        lines_list = list(lines)
        for idx, line in enumerate(lines_list):
            if idx == 0:
                bullets_tf.text = line
            else:
                p = bullets_tf.add_paragraph()
                p.text = line
            bullets_tf.paragraphs[idx].font.size = Pt(18)
        return top + Inches(0.6 + 0.4 * max(1, len(lines_list)))


__all__ = ["PPTXRenderer"]
